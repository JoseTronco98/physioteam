"""Exports that plug into the therapist's existing tools: Word documents and calendar files."""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor

from . import store


def document_docx(c: dict, doc: dict) -> bytes:
    v = doc["versions"][doc["version"] - 1]
    p = c["patients"][doc["patient_id"]]
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    h = d.add_heading(f"{doc['title']}", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0xC2, 0x18, 0x5B)
    d.add_paragraph(f"{c['clinic']['name']} · {c['clinic']['therapist']}")
    d.add_paragraph(f"Patient: {p['name']}  ·  Age: {p['age']}  ·  Version {v['version']} ({v['status']})  ·  {store.nice(v['created'])}")
    if v.get("fields"):
        t = d.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for f in v["fields"]:
            row = t.add_row().cells
            row[0].text = f["label"]
            row[1].text = f["value"] or "(empty - to be completed by the therapist)"
    if v.get("body_md"):
        for line in v["body_md"].splitlines():
            d.add_paragraph(line.replace("**", "").replace("_", ""))
    d.add_paragraph("")
    d.add_paragraph("Therapist signature: ______________________     Date: ____________")
    d.add_paragraph("Draft prepared with PhysioTeam AI. Reviewed and signed by the therapist.").runs[0].italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def document_html(c: dict, doc: dict) -> str:
    v = doc["versions"][doc["version"] - 1]
    p = c["patients"][doc["patient_id"]]
    rows = "".join(f"<tr><th>{f['label']}</th><td>{f['value'] or '<i>(empty)</i>'}</td></tr>" for f in v.get("fields") or [])
    body = (v.get("body_md") or "").replace("\n", "<br>")
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{doc['title']} - {p['name']}</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;max-width:760px;margin:30px auto;color:#111}}h1{{color:#c2185b}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:6px;text-align:left;vertical-align:top;font-size:13px}}
th{{width:34%;background:#fdeef5}}</style></head><body onload="window.print()">
<h1>{doc['title']}</h1><p>{c['clinic']['name']} · {c['clinic']['therapist']}<br>Patient: {p['name']}, {p['age']} · v{v['version']} ({v['status']})</p>
<table>{rows}</table><p>{body}</p><p><br>Therapist signature: ______________________</p></body></html>"""


def agenda_ics(c: dict) -> str:
    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//PhysioTeam//EN", "CALSCALE:GREGORIAN"]
    for a in sorted(c["appointments"].values(), key=lambda x: x["start"]):
        if a["status"] not in ("booked", "confirmed", "pending_approval"):
            continue
        p = c["patients"][a["patient_id"]]
        s = store.dt(a["start"]).strftime("%Y%m%dT%H%M%S")
        e = store.dt(a["end"]).strftime("%Y%m%dT%H%M%S")
        lines += ["BEGIN:VEVENT", f"UID:{a['id']}@physioteam", f"DTSTAMP:{stamp}", f"DTSTART:{s}", f"DTEND:{e}",
                  f"SUMMARY:{p['name']} - {a['service']}", f"DESCRIPTION:{p['condition']}", "END:VEVENT"]
    lines.append("END:VCALENDAR")
    return "\r\n".join(lines)
